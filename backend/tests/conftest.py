"""Builders for small course files, so tests never depend on private courses."""

import io

import pytest
from fpdf import FPDF
from pptx import Presentation


def _make_pdf(pages: list[str]) -> bytes:
    """One PDF page per string; an empty string gives a page without text."""
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 10, text)
    return bytes(pdf.output())


def _make_pptx(slides: list[list[str]]) -> bytes:
    """One slide per list: the first string is the title, the rest the body."""
    presentation = Presentation()
    for title, *body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "\n".join(body)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def make_pdf():
    return _make_pdf


@pytest.fixture
def make_pptx():
    return _make_pptx
