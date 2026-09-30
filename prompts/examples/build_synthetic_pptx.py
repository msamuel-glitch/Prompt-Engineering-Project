"""Build the PowerPoint version of synthetic courses from their Markdown files.

The PowerPoint files let anyone try the application on the synthetic courses
(S1, S2). Each Markdown file holds its slides in a ```text block, one
[Slide N] marker per slide. Run the script with the backend environment, which
includes python-pptx, and name the courses to build (all of them by default):

    backend/.venv/Scripts/python prompts/examples/build_synthetic_pptx.py injection_course_marketing_fr.md
    backend/.venv/bin/python prompts/examples/build_synthetic_pptx.py        (macOS/Linux, all courses)

Rebuilding a course changes its .pptx bytes even when the text is the same, so
only rebuild the courses whose Markdown changed.
"""

import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).parent
TITLE_AND_CONTENT, TITLE_ONLY = 1, 5  # layouts of the default template


def read_slides(source: Path) -> list[list[str]] | None:
    """Lines of each slide, or None if the file holds no slide course."""
    document = source.read_text(encoding="utf-8").replace("\r\n", "\n")
    course = re.search(r"^```text\n(\[Slide 1\].*?)^```", document, re.DOTALL | re.MULTILINE)
    if course is None:
        return None
    slides = re.split(r"^\[Slide \d+\]\n", course.group(1), flags=re.MULTILINE)[1:]
    return [slide.strip().splitlines() for slide in slides]


def build(source: Path) -> None:
    slides = read_slides(source)
    if slides is None:
        print(f"Skipped {source.name}: no [Slide 1] text block")
        return
    presentation = Presentation()
    for title, *body in slides:
        layout = presentation.slide_layouts[TITLE_AND_CONTENT if body else TITLE_ONLY]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        if body:
            slide.placeholders[1].text = "\n".join(body)
    target = source.with_suffix(".pptx")
    presentation.save(target)
    print(f"Wrote {target.name} ({len(presentation.slides)} slides)")


def main() -> None:
    names = sys.argv[1:]
    sources = [HERE / name for name in names] if names else sorted(HERE.glob("*.md"))
    for source in sources:
        build(source)


if __name__ == "__main__":
    main()
