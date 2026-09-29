"""Read the text of a course file, page by page or slide by slide."""

import io
import logging
from dataclasses import dataclass
from pathlib import PurePath

import pdfplumber
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas import SourceType

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """The file cannot be turned into course text; the message says why."""


class UnsupportedFileError(ExtractionError):
    pass


@dataclass
class CourseUnit:
    """One page of a PDF or one slide of a presentation."""

    number: int  # starts at 1, as displayed by PDF readers and PowerPoint
    text: str


@dataclass
class Course:
    source_type: SourceType
    units: list[CourseUnit]

    @property
    def empty_units(self) -> list[int]:
        return [unit.number for unit in self.units if not unit.text]

    def to_prompt_text(self) -> str:
        """Course text with a [Page N] or [Slide N] marker before each unit.

        Units without text are left out: there is nothing in them to cite.
        """
        label = self.source_type.capitalize()
        return "\n\n".join(
            f"[{label} {unit.number}]\n{unit.text}" for unit in self.units if unit.text
        )


def extract_course(file_name: str, data: bytes) -> Course:
    suffix = PurePath(file_name).suffix.lower()
    if suffix == ".pdf":
        course = _extract_pdf(data)
    elif suffix == ".pptx":
        course = _extract_pptx(data)
    else:
        raise UnsupportedFileError("Only PDF and PPTX files are supported.")

    if not any(unit.text for unit in course.units):
        raise ExtractionError(
            "No text could be extracted. The file may contain only images, "
            "like a scanned PDF; text recognition (OCR) is not supported yet."
        )
    return course


def _extract_pdf(data: bytes) -> Course:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            units = [
                CourseUnit(number, _clean(page.extract_text() or ""))
                for number, page in enumerate(pdf.pages, start=1)
            ]
    # Malformed files raise many different exception types inside pdfminer.
    except Exception as error:
        logger.warning("PDF extraction failed", exc_info=True)
        raise ExtractionError(
            "This PDF could not be read. It may be damaged or password-protected."
        ) from error
    return Course("page", units)


def _extract_pptx(data: bytes) -> Course:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as error:
        logger.warning("PPTX extraction failed", exc_info=True)
        raise ExtractionError(
            "This PowerPoint file could not be read. It may be damaged."
        ) from error
    units = [
        CourseUnit(number, _clean("\n".join(_shape_texts(slide.shapes))))
        for number, slide in enumerate(presentation.slides, start=1)
    ]
    return Course("slide", units)


def _shape_texts(shapes) -> list[str]:
    texts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_shape_texts(shape.shapes))
        elif shape.has_text_frame:
            texts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                texts.append(" | ".join(cell.text for cell in row.cells))
    return texts


def _clean(text: str) -> str:
    lines = (line.strip() for line in text.splitlines())
    return "\n".join(line for line in lines if line)
